"""Regresiones del flujo operativo y de la adaptacion multiusuario, con datos sinteticos."""

from __future__ import annotations

import io
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from threading import BoundedSemaphore
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl

import validacion_materiales_prestaciones_AGP as motor
from catalogo_aportacion import cargar_catalogo, serializar_catalogo
from scripts.actualizar_maestro import actualizar, leer_entrada
from validacion_web import ServidorOcupado, reservar_proceso, validar_archivos, verificar_excel


def libro_bytes(wb):
    salida = io.BytesIO()
    wb.save(salida)
    wb.close()
    return salida.getvalue()


def pms_rg11(filas):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RG11_budget"
    ws.append(["PMS de prueba"])
    ws.append(["cod_prest/mat", "quantity", "descripcion"])
    for fila in filas:
        ws.append(fila)
    return libro_bytes(wb)


def pms_e1():
    wb = openpyxl.Workbook()
    ser = wb.active
    ser.title = "RG10_SER"
    ser.append(["PMS E1 sintetico"])
    ser.append([])
    ser.append(["PRESTACION", "CANTIDAD", "DESCRIPCION"])
    ser.append(["P001", 1, "Servicio"])
    mat = wb.create_sheet("RG36_MAT")
    mat.append([])
    mat.append([])
    mat.append(["CODIGO MATERIAL E4E", "CANTIDAD", "TIPO DE MOVIMIENTO", "DESCRIPCION"])
    mat.append(["120570", 1, "Retirado", "Transformador"])
    mat.append(["970698", -2, "Retirado", "Aisladores"])
    mat.append(["001234", 4, "Material de aportacion", "Aportacion"])
    return libro_bytes(wb)


def agp_bytes(proyectos, historico=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.append(["Resumen sintetico"])
    for codigo, protocolos in proyectos:
        ws.append([f"{codigo} - Proyecto sintetico"])
        for protocolo, filas in protocolos:
            ws.append([f"JA01 - {protocolo} - Protocolo sintetico"])
            ws.row_dimensions[ws.max_row].outlineLevel = 1
            ws.append(["UT de prueba"])
            ws.row_dimensions[ws.max_row].outlineLevel = 2
            ws.append(["CIT de prueba"])
            ws.row_dimensions[ws.max_row].outlineLevel = 3
            for texto in filas:
                ws.append([texto])
                ws.row_dimensions[ws.max_row].outlineLevel = 5 if historico else 4
    return libro_bytes(wb)


class PilotoTest(unittest.TestCase):
    def setUp(self):
        self.temporal = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporal.cleanup)
        self.raiz = Path(self.temporal.name)
        self.ruta_catalogo = self.raiz / "catalogo.csv"
        self.ruta_catalogo.write_bytes(serializar_catalogo(["001234"]))
        self.catalogo = cargar_catalogo(self.ruta_catalogo)

    def validar(self, pms, agp, catalogo=None):
        return validar_archivos("PMS.xlsx", pms, "data.xlsx", agp, catalogo or self.catalogo, BoundedSemaphore(1))

    def test_e1_retirados_negativos_y_aportacion(self):
        agp = agp_bytes([("2000138100", [("PR1", [
            "P001 - T0 - Servicio - 1.000 - 001",
            "120570 - TRAFO 1F - Rev - 1.000- - 001",
            "970698 - AISLADORES - 2.000- - 001",
        ])])])
        resultado = self.validar(pms_e1(), agp)
        self.assertEqual(resultado["datos_resumen"]["SEMAFORO"], "OK")
        self.assertEqual(resultado["datos_resumen"]["TOTAL_APORTACION_EXCLUIDOS"], 1)
        tabla = resultado["comparacion"].set_index("CODIGO")
        self.assertEqual(tabla.loc["120570", "CANTIDAD_PMS"], -1)
        self.assertEqual(tabla.loc["970698", "CANTIDAD_PMS"], -2)
        self.assertEqual(resultado["datos_resumen"]["VERSION_MAESTRO_SHA256"], "No aplica: E1")

    def test_varios_agp_protocolos_y_nombre_salida(self):
        pms = pms_rg11([("456", 3, "Material"), ("P001", 3, "Servicio"), ("001234", 2, "Aportacion")])
        agp = agp_bytes([
            ("2000139288", [("PR1", ["456 - Material - 1.000 - 001"]),
                            ("PR2", ["P001 - T0 - Servicio - 3.000 - 001"])]),
            ("2000139289", [("PR3", ["456 - Material - 2.000 - 001"])]),
        ])
        resultado = self.validar(pms, agp)
        self.assertEqual(resultado["datos_resumen"]["SEMAFORO"], "OK")
        self.assertEqual(resultado["contexto_agp"]["TOTAL_AGPS"], 2)
        self.assertEqual(resultado["contexto_agp"]["TOTAL_PROTOCOLOS"], 3)
        self.assertIn("2000139288_2000139289", resultado["nombre_excel"])
        contexto = resultado["agp_normalizado"].set_index("CODIGO").loc["456", "AGP"]
        self.assertIn("2000139288", contexto)
        self.assertIn("2000139289", contexto)
        wb = openpyxl.load_workbook(io.BytesIO(resultado["archivo_excel"]))
        self.addCleanup(wb.close)
        self.assertEqual(len(wb.sheetnames), 10)
        self.assertTrue(all(len(ws._images) == 1 for ws in wb.worksheets))
        self.assertEqual(wb["Comparacion"].freeze_panes, "A6")
        self.assertTrue(wb["Comparacion"].auto_filter.ref)
        resumen = {r[0].value: r[1].value for r in wb["Resumen"].iter_rows(min_row=8)}
        self.assertEqual(resumen["VERSION_MAESTRO_SHA256"], self.catalogo.version)
        self.assertEqual(resumen["ZONA_HORARIA"], "America/Bogota")

    def test_aportacion_valorizada_en_agp_no_se_oculta(self):
        resultado = self.validar(
            pms_rg11([("001234", 1, "Aportacion"), ("456", 1, "Normal")]),
            agp_bytes([("2000139288", [("PR1", ["001234 - Aportacion - 1 - 001", "456 - Normal - 1 - 001"])])]),
        )
        self.assertEqual(resultado["comparacion"].set_index("CODIGO").loc["001234", "ESTADO_VALIDACION"], "SOLO_EN_AGP")

    def test_diferencia_solos_y_error_cantidad_generan_informe(self):
        resultado = self.validar(
            pms_rg11([("456", "1,5", "Material"), ("777", 1, "Solo PMS"), ("P001", "invalida", "Servicio")]),
            agp_bytes([("2000139288", [("PR1", ["456 - Material - 2 - 001", "888 - Solo AGP - 1 - 001", "P001 - T0 - Servicio - 1 - 001"])])]),
        )
        estados = resultado["comparacion"].set_index("CODIGO")["ESTADO_VALIDACION"]
        self.assertEqual(estados["456"], "DIFERENCIA_CANTIDAD")
        self.assertEqual(estados["777"], "SOLO_EN_PMS")
        self.assertEqual(estados["888"], "SOLO_EN_AGP")
        self.assertEqual(estados["P001"], "ERROR_DATO")
        self.assertEqual(resultado["datos_resumen"]["SEMAFORO"], "ERROR")
        self.assertGreater(len(resultado["archivo_excel"]), 0)

    def test_estructura_historica(self):
        resultado = self.validar(pms_rg11([("456", 1, "Material")]), agp_bytes([
            ("2000139288", [("PR1", ["456 - Material - 1 - 001"])])], historico=True))
        self.assertEqual(resultado["datos_resumen"]["SEMAFORO"], "OK")
        self.assertEqual(resultado["contexto_agp"]["NIVELES_CONCEPTUALES"], 6)

    def test_sesiones_mismos_nombres_no_comparten_archivos_y_limpian_temporales(self):
        rutas = []
        original = motor.ejecutar_validacion

        def observar(*args, **kwargs):
            rutas.append(Path(args[0]).parents[1])
            return original(*args, **kwargs)

        def caso(cantidad):
            return self.validar(pms_rg11([("456", cantidad, "Material")]), agp_bytes([
                ("2000139288", [("PR1", [f"456 - Material - {cantidad} - 001"])])]))

        with patch.object(motor, "ejecutar_validacion", side_effect=observar):
            with ThreadPoolExecutor(max_workers=2) as pool:
                resultados = list(pool.map(caso, [1, 7]))
        self.assertNotEqual(rutas[0], rutas[1])
        self.assertTrue(all(not ruta.exists() for ruta in rutas))
        self.assertEqual([r["comparacion"].iloc[0]["CANTIDAD_PMS"] for r in resultados], [1, 7])
        self.assertTrue(all(r["datos_resumen"]["SEMAFORO"] == "OK" for r in resultados))
        self.assertIsNone(motor._DIRECTORIO_TEMPORAL.get())

    def test_servidor_ocupado_y_liberacion_por_error(self):
        semaforo = BoundedSemaphore(1)
        with reservar_proceso(semaforo):
            with self.assertRaises(ServidorOcupado):
                with reservar_proceso(semaforo):
                    self.fail("No debe entrar cuando esta ocupado.")
        with self.assertRaises(RuntimeError):
            with reservar_proceso(semaforo):
                raise RuntimeError("Fallo simulado")
        with reservar_proceso(semaforo):
            pass

    def test_excel_invalido_y_dimension_excesiva(self):
        with self.assertRaisesRegex(ValueError, "No se pudo abrir"):
            verificar_excel(b"No soy Excel", "PMS")
        wb = openpyxl.Workbook()
        wb.active["A50001"] = "Dato"
        with self.assertRaisesRegex(ValueError, "demasiado extensa"):
            verificar_excel(libro_bytes(wb), "PMS")

    def test_maestro_vacio_bloquea_rg11_permite_e1(self):
        self.ruta_catalogo.write_bytes(serializar_catalogo([]))
        vacio = cargar_catalogo(self.ruta_catalogo)
        with self.assertRaisesRegex(ValueError, "maestro esta vacio"):
            self.validar(pms_rg11([("456", 1, "Material")]), agp_bytes([
                ("2000139288", [("PR1", ["456 - Material - 1 - 001"])])]), vacio)
        resultado = validar_archivos("E1.xlsx", pms_e1(), "data.xlsx", agp_bytes([
            ("2000138100", [("PR1", ["P001 - T0 - Servicio - 1 - 001",
                                    "120570 - TRAFO - 1.000- - 001", "970698 - AISLADORES - 2.000- - 001"])])
        ]), None, BoundedSemaphore(1))
        self.assertEqual(resultado["datos_resumen"]["SEMAFORO"], "OK")

    def test_pms_sin_hojas_requeridas_exporta_error_de_estructura(self):
        wb = openpyxl.Workbook()
        wb.active["A1"] = "Estructura incorrecta"
        resultado = self.validar(libro_bytes(wb), agp_bytes([
            ("2000139288", [("PR1", ["456 - Material - 1 - 001"])])]))
        self.assertEqual(resultado["datos_resumen"]["SEMAFORO"], "ERROR")
        self.assertFalse(resultado["errores"].empty)
        self.assertIn("No se pudo detectar", resultado["errores"].iloc[0]["MENSAJE"])
        self.assertTrue(resultado["archivo_excel"])

    def test_catalogo_masivo_duplicados_ceros_y_vista_previa(self):
        original = self.ruta_catalogo.read_bytes()
        informe = actualizar("agregar", ["001234", "456", "456"], self.ruta_catalogo)
        self.assertEqual(informe["agregados"], 1)
        self.assertEqual(informe["repetidos_en_carga"], 1)
        self.assertEqual(informe["ya_existentes"], 1)
        self.assertEqual(self.ruta_catalogo.read_bytes(), original)
        actualizar("agregar", ["001234", "456", "456"], self.ruta_catalogo, True)
        self.assertEqual(cargar_catalogo(self.ruta_catalogo).codigos, ("001234", "456"))
        self.assertTrue((self.raiz / ".respaldos" / f"{self.catalogo.version}.csv").is_file())
        with self.assertRaises(ValueError):
            actualizar("agregar", ["999", "abc"], self.ruta_catalogo, True)
        self.assertNotIn("999", cargar_catalogo(self.ruta_catalogo).codigos)
        actualizar("retirar", ["456", "999"], self.ruta_catalogo, True)
        self.assertEqual(cargar_catalogo(self.ruta_catalogo).codigos, ("001234",))

    def test_archivos_csv_txt_excel_y_normalizacion_catalogo(self):
        csv = self.raiz / "entrada.csv"
        csv.write_text("CODIGO;DESCRIPCION\n001234;Material\n456;Otro\n", encoding="utf-8-sig")
        self.assertEqual(leer_entrada(csv), ["001234", "456"])
        txt = self.raiz / "entrada.txt"
        txt.write_text("001234\n456 456", encoding="utf-8")
        self.assertEqual(leer_entrada(txt), ["001234", "456", "456"])
        wb = openpyxl.Workbook()
        wb.active.append(["MAT APORT"])
        wb.active.append(["001234"])
        wb.active.append([456])
        xlsx = self.raiz / "entrada.xlsx"
        wb.save(xlsx)
        wb.close()
        self.assertEqual(leer_entrada(xlsx), ["001234", "456"])
        self.ruta_catalogo.write_text("CODIGO\n456\n456\n001234\n", encoding="utf-8")
        with self.assertRaises(ValueError):
            cargar_catalogo(self.ruta_catalogo)
        actualizar("normalizar", [], self.ruta_catalogo, True)
        self.assertEqual(cargar_catalogo(self.ruta_catalogo).codigos, ("001234", "456"))

    def test_formulas_de_descripcion_se_exportan_como_texto(self):
        agp = agp_bytes([("2000139288", [("PR1", ["456 - =1+1 - 1 - 001"])])])
        resultado = self.validar(pms_rg11([("456", 1, "Material")]), agp)
        wb = openpyxl.load_workbook(io.BytesIO(resultado["archivo_excel"]), data_only=False)
        self.addCleanup(wb.close)
        self.assertFalse(any(c.data_type == "f" for ws in wb for fila in ws for c in fila))

    def test_copia_saneada_se_crea_en_temporal_de_la_sesion(self):
        original = self.raiz / "original.xlsx"
        original.write_bytes(pms_rg11([("456", 1, "Material")]))
        contenido_original = original.read_bytes()
        aislado = self.raiz / "sesion"
        aislado.mkdir()
        token = motor._DIRECTORIO_TEMPORAL.set(aislado)
        try:
            ruta = Path(motor._crear_copia_excel_sin_defined_names_corruptos(str(original)))
            self.assertEqual(ruta.parent, aislado)
            self.assertEqual(original.read_bytes(), contenido_original)
        finally:
            motor._DIRECTORIO_TEMPORAL.reset(token)


try:
    from streamlit.testing.v1 import AppTest
except ImportError:
    AppTest = None


@unittest.skipIf(AppTest is None, "Instale requirements.txt para comprobar la interfaz Streamlit.")
class InterfazTest(unittest.TestCase):
    def test_navegacion_y_reinicio_sin_errores(self):
        app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / "streamlit_app.py")).run(timeout=30)
        self.assertEqual(len(app.exception), 0)
        self.assertEqual(len(app.get("file_uploader")), 2)
        self.assertTrue(next(b for b in app.button if b.label.startswith("Validar materiales")).disabled)
        app.sidebar.radio[0].set_value("Manual de uso").run()
        self.assertEqual(len(app.exception), 0)
        self.assertTrue(any("Guia" in m.value or "Guía" in m.value for m in app.markdown))
        app.sidebar.radio[0].set_value("Maestro de aportación").run()
        self.assertEqual(len(app.exception), 0)
        self.assertEqual(len(app.dataframe), 1)
        app.text_input[0].set_value("imposible").run()
        self.assertEqual(len(app.exception), 0)
        app.sidebar.radio[0].set_value("Validar caso").run()
        next(b for b in app.button if b.label == "Nuevo caso").click().run()
        self.assertEqual(len(app.exception), 0)
        self.assertEqual(len(app.get("file_uploader")), 2)


if __name__ == "__main__":
    unittest.main()
