"""Administracion local del maestro: vista previa por defecto, escritura con --guardar."""

from __future__ import annotations

import argparse
import csv
import io
import os
import re
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl

from catalogo_aportacion import RUTA_CATALOGO, cargar_catalogo, serializar_catalogo, validar_codigo_material
from validacion_materiales_prestaciones_AGP import _cargar_workbook_seguro, normalizar_encabezado

ALIAS = ("MAT APORT", "CODIGO MATERIAL", "CODIGO", "MATERIAL", "UBICACION")


def leer_entrada(ruta: Path) -> list[str]:
    if ruta.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = _cargar_workbook_seguro(str(ruta), read_only=True, data_only=True, keep_links=False)
        try:
            candidatos = []
            for ws in wb.worksheets:
                for numero, fila in enumerate(ws.iter_rows(max_row=min(ws.max_row, 30), values_only=True), 1):
                    nombres = [normalizar_encabezado(v) for v in fila]
                    for prioridad, alias in enumerate(ALIAS):
                        if alias in nombres:
                            candidatos.append((prioridad, ws.title, numero, nombres.index(alias)))
            if not candidatos:
                raise ValueError("No se encontro una columna MAT APORT, CODIGO MATERIAL o CODIGO.")
            _, hoja, cabecera, columna = min(candidatos)
            valores = [fila[columna] for fila in wb[hoja].iter_rows(min_row=cabecera + 1, values_only=True)
                       if len(fila) > columna and fila[columna] is not None and str(fila[columna]).strip()]
        finally:
            wb.close()
    elif ruta.suffix.lower() == ".csv":
        texto = ruta.read_text(encoding="utf-8-sig")
        try:
            dialecto = csv.Sniffer().sniff(texto[:8192], delimiters=",;\t")
        except csv.Error:
            dialecto = csv.excel
        filas = list(csv.reader(io.StringIO(texto), dialecto))
        if not filas:
            raise ValueError("El archivo esta vacio.")
        nombres = [normalizar_encabezado(v) for v in filas[0]]
        columna = next((nombres.index(a) for a in ALIAS if a in nombres), None)
        if columna is None:
            raise ValueError("El CSV debe tener un encabezado CODIGO o MAT APORT.")
        valores = [fila[columna] for fila in filas[1:] if len(fila) > columna and fila[columna].strip()]
    elif ruta.suffix.lower() == ".txt":
        valores = re.split(r"[\s,;]+", ruta.read_text(encoding="utf-8-sig").strip())
    else:
        raise ValueError("La entrada debe ser Excel .xlsx/.xlsm, CSV o TXT.")
    if not valores:
        raise ValueError("No hay codigos para procesar.")
    return [validar_codigo_material(v) for v in valores]


def actualizar(accion: str, entrada: list[str], ruta: Path, guardar: bool = False) -> dict:
    actual = cargar_catalogo(ruta, permitir_duplicados=accion == "normalizar")
    existentes = set(actual.codigos)
    lista = [validar_codigo_material(c) for c in entrada]
    recibidos = set(lista)
    if accion == "importar":
        nuevos = recibidos
    elif accion == "agregar":
        nuevos = existentes | recibidos
    elif accion == "retirar":
        nuevos = existentes - recibidos
    elif accion == "normalizar":
        nuevos = existentes
    else:
        raise ValueError("Accion no reconocida.")
    informe = {
        "recibidos": len(lista), "repetidos_en_carga": len(lista) - len(recibidos),
        "ya_existentes": len(existentes & recibidos), "agregados": len(nuevos - existentes),
        "retirados": len(existentes - nuevos), "total_final": len(nuevos), "guardado": False,
    }
    if guardar:
        contenido = serializar_catalogo(nuevos)
        if contenido != actual.contenido:
            respaldos = ruta.parent / ".respaldos"
            respaldos.mkdir(exist_ok=True)
            (respaldos / f"{actual.version}.csv").write_bytes(actual.contenido)
            with tempfile.NamedTemporaryFile(dir=ruta.parent, suffix=".tmp", delete=False) as salida:
                temporal = Path(salida.name)
                salida.write(contenido)
            try:
                os.replace(temporal, ruta)
            finally:
                temporal.unlink(missing_ok=True)
        informe["guardado"] = True
    return informe


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("accion", choices=["listar", "importar", "agregar", "retirar", "normalizar"])
    origen = parser.add_mutually_exclusive_group()
    origen.add_argument("--archivo", type=Path)
    origen.add_argument("--codigos", help="Codigos separados por espacios, comas o punto y coma.")
    parser.add_argument("--catalogo", type=Path, default=RUTA_CATALOGO)
    parser.add_argument("--guardar", action="store_true", help="Aplicar cambios. Sin esta opcion solo muestra una vista previa.")
    args = parser.parse_args()
    try:
        if args.accion == "listar":
            catalogo = cargar_catalogo(args.catalogo)
            print(f"Total: {len(catalogo.codigos)} | SHA-256: {catalogo.version}")
            print("\n".join(catalogo.codigos))
            return 0
        if args.accion != "normalizar" and not (args.archivo or args.codigos):
            parser.error("Indique --archivo o --codigos.")
        entrada = leer_entrada(args.archivo) if args.archivo else [
            validar_codigo_material(c) for c in re.split(r"[\s,;]+", (args.codigos or "").strip()) if c
        ]
        informe = actualizar(args.accion, entrada, args.catalogo, args.guardar)
        for clave, valor in informe.items():
            print(f"{clave}: {valor}")
        if not args.guardar:
            print("VISTA PREVIA. Repita el comando con --guardar para aplicar los cambios.")
        else:
            print("Maestro guardado localmente. Publique el CSV actualizado en GitHub para actualizar la web.")
        return 0
    except (ValueError, OSError, openpyxl.utils.exceptions.InvalidFileException) as exc:
        print(f"No se modifico el maestro: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
